"""
Tests for: python manage.py bm_export <run_id> [--format csv|excel] [--output path]
"""
import csv
import os
import tempfile

from django.core.management.base import CommandError

from apps.runs.models import BenchmarkRun, RunResult
from tests.base import BaseCommandTest, make_run


class TestBmExport(BaseCommandTest):

    def setUp(self):
        super().setUp()
        self.run = make_run(self.benchmark, self.provider)
        # Add a result
        question = self.benchmark.questions.first()
        RunResult.objects.create(
            run=self.run,
            question=question,
            model_response='A',
            parsed_answer='A',
            is_correct=True,
            response_time=0.5,
            tokens_input=100,
            tokens_output=5,
            estimated_cost=0.001,
        )

    def test_export_csv_default(self):
        with tempfile.NamedTemporaryFile(suffix='.csv', delete=False) as f:
            path = f.name
        try:
            out = self.call('bm_export', str(self.run.id), '--output', path)
            self.assertIn('Exported', out)
            self.assertTrue(os.path.exists(path))
            with open(path, newline='', encoding='utf-8') as f:
                reader = csv.reader(f)
                rows = list(reader)
            # Header + 1 data row
            self.assertEqual(len(rows), 2)
            self.assertIn('question_id', rows[0])
            self.assertIn('is_correct', rows[0])
        finally:
            if os.path.exists(path):
                os.unlink(path)

    def test_export_csv_headers(self):
        with tempfile.NamedTemporaryFile(suffix='.csv', delete=False) as f:
            path = f.name
        try:
            self.call('bm_export', str(self.run.id), '--output', path)
            with open(path, newline='', encoding='utf-8') as f:
                headers = next(csv.reader(f))
            expected = ['question_id', 'subject', 'difficulty', 'question',
                        'correct_answer', 'model_answer', 'parsed_answer', 'is_correct',
                        'response_time_s', 'tokens_input', 'tokens_output', 'cost_usd']
            self.assertEqual(headers, expected)
        finally:
            if os.path.exists(path):
                os.unlink(path)

    def test_export_excel(self):
        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as f:
            path = f.name
        try:
            out = self.call('bm_export', str(self.run.id),
                            '--format', 'excel', '--output', path)
            self.assertIn('Exported', out)
            self.assertTrue(os.path.exists(path))
            import openpyxl
            wb = openpyxl.load_workbook(path)
            self.assertIn('Summary', wb.sheetnames)
            self.assertIn('By Subject', wb.sheetnames)
            self.assertIn('Results', wb.sheetnames)
        finally:
            if os.path.exists(path):
                os.unlink(path)

    def test_export_run_not_found_raises(self):
        with self.assertRaises(CommandError):
            self.call('bm_export', '99999')

    def test_export_default_filename(self):
        """Without --output, file should be created as run_<id>.csv in cwd."""
        expected = f'run_{self.run.id}.csv'
        try:
            out = self.call('bm_export', str(self.run.id))
            self.assertIn('Exported', out)
            self.assertTrue(os.path.exists(expected))
        finally:
            if os.path.exists(expected):
                os.unlink(expected)

    def test_export_empty_run_still_works(self):
        """Export should succeed even if the run has no results."""
        run2 = make_run(self.benchmark, self.provider, model_name='empty-model')
        with tempfile.NamedTemporaryFile(suffix='.csv', delete=False) as f:
            path = f.name
        try:
            out = self.call('bm_export', str(run2.id), '--output', path)
            self.assertIn('Exported', out)
            with open(path, newline='', encoding='utf-8') as f:
                rows = list(csv.reader(f))
            self.assertEqual(len(rows), 1)  # Header only
        finally:
            if os.path.exists(path):
                os.unlink(path)
