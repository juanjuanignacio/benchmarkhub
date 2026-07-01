import csv
import io
import logging
import math
import threading

from django.contrib import messages
from django.core.paginator import Paginator
from django.db.models import Avg, Max, Count, F, Q, ExpressionWrapper, FloatField
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.views.generic import DetailView, ListView

from .models import Benchmark, BenchmarkQuestion, PromptTemplate


def _wilson_ci(correct: int, total: int) -> tuple:
    """95% Wilson score confidence interval. Returns (lower_pct, upper_pct)."""
    if total == 0:
        return None, None
    z = 1.960
    p_hat = correct / total
    z2 = z * z
    denom = 1 + z2 / total
    centre = (p_hat + z2 / (2 * total)) / denom
    margin = (z / denom) * math.sqrt(p_hat * (1 - p_hat) / total + z2 / (4 * total * total))
    return round(max(0.0, (centre - margin) * 100), 1), round(min(100.0, (centre + margin) * 100), 1)
from .registry import BENCHMARK_REGISTRY, get_loader

logger = logging.getLogger(__name__)


class BenchmarkListView(ListView):
    template_name = 'benchmarks/list.html'
    context_object_name = 'benchmarks'

    def get_queryset(self):
        category = self.request.GET.get('category', '')
        btype = self.request.GET.get('type', '')
        qs = Benchmark.objects.all()
        if category:
            qs = qs.filter(category=category)
        if btype:
            qs = qs.filter(benchmark_type=btype)
        return qs

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        loaded_slugs = set(Benchmark.objects.filter(loaded_at__isnull=False).values_list('slug', flat=True))
        registry_info = []
        for slug, loader_cls in BENCHMARK_REGISTRY.items():
            loader = loader_cls()
            try:
                bench = Benchmark.objects.get(slug=slug)
            except Benchmark.DoesNotExist:
                bench = None
            registry_info.append({
                'slug': slug,
                'name': loader.name,
                'description': loader.description,
                'category': loader.category,
                'benchmark_type': getattr(loader, 'benchmark_type', 'text'),
                'benchmark': bench,
                'is_loaded': bench is not None and bench.is_loaded,
            })
        ctx['registry_info'] = registry_info
        ctx['category_filter'] = self.request.GET.get('category', '')
        ctx['type_filter'] = self.request.GET.get('type', '')
        ctx['categories'] = Benchmark.CATEGORY_CHOICES
        ctx['benchmark_types'] = Benchmark.BENCHMARK_TYPE_CHOICES
        return ctx


class BenchmarkDetailView(DetailView):
    model = Benchmark
    template_name = 'benchmarks/detail.html'
    context_object_name = 'benchmark'

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        benchmark = self.object
        ctx['runs'] = benchmark.runs.select_related('provider').order_by('-created_at')[:10]
        ctx['question_count'] = benchmark.questions.count()
        ctx['subjects'] = (
            benchmark.questions.exclude(subject='')
            .values_list('subject', flat=True)
            .distinct()[:20]
        )
        # Prompt template context
        loader = get_loader(benchmark.slug)
        ctx['loader'] = loader
        if loader:
            sample_q = benchmark.questions.first()
            if sample_q:
                try:
                    ctx['default_prompt_preview'] = loader.format_prompt(sample_q)
                except Exception:
                    ctx['default_prompt_preview'] = None
            else:
                ctx['default_prompt_preview'] = None
        return ctx


class BenchmarkQuestionsView(ListView):
    template_name = 'benchmarks/questions.html'
    context_object_name = 'questions'
    paginate_by = 20

    def get_queryset(self):
        self.benchmark = get_object_or_404(Benchmark, slug=self.kwargs['slug'])
        qs = BenchmarkQuestion.objects.filter(benchmark=self.benchmark)
        search = self.request.GET.get('search', '')
        if search:
            qs = qs.filter(question__icontains=search)
        subject = self.request.GET.get('subject', '')
        if subject:
            qs = qs.filter(subject=subject)
        return qs

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        ctx['benchmark'] = self.benchmark
        ctx['search'] = self.request.GET.get('search', '')
        ctx['subject_filter'] = self.request.GET.get('subject', '')
        ctx['subjects'] = (
            BenchmarkQuestion.objects.filter(benchmark=self.benchmark)
            .exclude(subject='')
            .values_list('subject', flat=True)
            .distinct()
            .order_by('subject')
        )
        return ctx


def _do_load_benchmark(slug, num_samples, benchmark_id):
    """Background function to load benchmark questions."""
    import django
    django.setup()
    from apps.benchmarks.models import Benchmark, BenchmarkQuestion

    try:
        benchmark = Benchmark.objects.get(id=benchmark_id)
        loader = get_loader(slug)
        if not loader:
            benchmark.metadata = {'error': f'No loader for slug: {slug}'}
            benchmark.save()
            return

        # Mark as loading with progress tracking
        benchmark.metadata = {'load_status': 'downloading', 'load_progress': 0, 'load_total': 0}
        benchmark.save()

        questions_data = loader.load_questions()
        if num_samples and num_samples > 0:
            questions_data = questions_data[:num_samples]

        total = len(questions_data)
        benchmark.metadata = {'load_status': 'saving', 'load_progress': 0, 'load_total': total}
        benchmark.save()

        # Delete old questions in its own short transaction
        BenchmarkQuestion.objects.filter(benchmark=benchmark).delete()

        # Insert in small batches WITHOUT wrapping everything in one giant
        # transaction.  Each bulk_create grabs the SQLite write-lock briefly
        # and releases it, so the web server is never blocked for long.
        BATCH = 200
        saved = 0
        batch = []
        for qdata in questions_data:
            batch.append(BenchmarkQuestion(
                benchmark=benchmark,
                question_id=str(qdata['question_id']),
                question=qdata['question'],
                choice_a=qdata.get('choice_a'),
                choice_b=qdata.get('choice_b'),
                choice_c=qdata.get('choice_c'),
                choice_d=qdata.get('choice_d'),
                correct_answer=str(qdata.get('correct_answer', '')),
                subject=str(qdata.get('subject', '')),
                difficulty=str(qdata.get('difficulty', '')),
                context=str(qdata.get('context', '')),
                image_paths=qdata.get('image_paths', []),
                audio_path=str(qdata.get('audio_path', '')),
                metadata=qdata.get('metadata', {}),
            ))
            if len(batch) >= BATCH:
                BenchmarkQuestion.objects.bulk_create(batch, ignore_conflicts=True)
                saved += len(batch)
                batch = []
                # Update progress so the UI can poll it
                benchmark.metadata = {'load_status': 'saving', 'load_progress': saved, 'load_total': total}
                benchmark.save(update_fields=['metadata'])
        if batch:
            BenchmarkQuestion.objects.bulk_create(batch, ignore_conflicts=True)
            saved += len(batch)

        benchmark.num_questions = BenchmarkQuestion.objects.filter(benchmark=benchmark).count()
        benchmark.loaded_at = timezone.now()
        meta = {'load_status': 'done'}
        if getattr(loader, 'excludes_images', False):
            meta['excludes_images'] = True
        benchmark.metadata = meta
        benchmark.save()
        logger.info(f"Loaded {benchmark.num_questions} questions for {slug}")

    except Exception as e:
        logger.error(f"Error loading benchmark {slug}: {e}", exc_info=True)
        try:
            benchmark = Benchmark.objects.get(id=benchmark_id)
            error_str = str(e)
            if 'gated dataset' in error_str.lower():
                error_type = 'gated'
            elif 'NameResolutionError' in error_str or 'ConnectionError' in error_str or 'Max retries' in error_str:
                error_type = 'network'
                error_str = (
                    'Network error: could not connect to HuggingFace CDN. '
                    'Check your internet connection and DNS settings, then try again.'
                )
            else:
                error_type = 'general'
            benchmark.metadata = {'error': error_str, 'error_type': error_type}
            benchmark.save()
        except Exception:
            pass


def load_benchmark_view(request, slug):
    if request.method != 'POST':
        return redirect('benchmarks:list')

    num_samples = int(request.POST.get('num_samples', 0) or 0)
    loader = get_loader(slug)
    if not loader:
        messages.error(request, f'Unknown benchmark: {slug}')
        return redirect('benchmarks:list')

    benchmark_type = getattr(loader, 'benchmark_type', 'text')
    benchmark, created = Benchmark.objects.get_or_create(
        slug=slug,
        defaults={
            'name': loader.name,
            'description': loader.description,
            'category': loader.category,
            'benchmark_type': benchmark_type,
        }
    )
    if not created:
        benchmark.name = loader.name
        benchmark.description = loader.description
        benchmark.category = loader.category
        benchmark.benchmark_type = benchmark_type
        benchmark.save()

    t = threading.Thread(
        target=_do_load_benchmark,
        args=(slug, num_samples, benchmark.id),
        daemon=True,
    )
    t.start()

    sample_msg = f' (first {num_samples} samples)' if num_samples else ''
    messages.success(
        request,
        f'Loading {loader.name}{sample_msg} in background. Refresh the page in a few moments.'
    )
    return redirect('benchmarks:list')


def benchmark_load_status(request, slug):
    """AJAX endpoint: returns loading progress for a benchmark."""
    try:
        benchmark = Benchmark.objects.get(slug=slug)
    except Benchmark.DoesNotExist:
        return JsonResponse({'slug': slug, 'status': 'idle', 'progress': 0, 'total': 0,
                             'error': '', 'error_type': '', 'num_questions': 0, 'is_loaded': False})
    meta = benchmark.metadata or {}
    return JsonResponse({
        'slug': slug,
        'status': meta.get('load_status', 'idle'),
        'progress': meta.get('load_progress', 0),
        'total': meta.get('load_total', 0),
        'error': meta.get('error', ''),
        'error_type': meta.get('error_type', ''),
        'num_questions': benchmark.num_questions,
        'is_loaded': benchmark.is_loaded,
    })


def delete_benchmark_view(request, slug):
    if request.method != 'POST':
        return redirect('benchmarks:list')
    benchmark = get_object_or_404(Benchmark, slug=slug)
    name = benchmark.name
    benchmark.delete()
    messages.success(request, f'Benchmark "{name}" deleted.')
    return redirect('benchmarks:list')


def export_questions_csv(request, slug):
    benchmark = get_object_or_404(Benchmark, slug=slug)
    questions = BenchmarkQuestion.objects.filter(benchmark=benchmark)

    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = f'attachment; filename="{slug}_questions.csv"'

    is_rag = (benchmark.metadata or {}).get('is_rag', False)
    writer = csv.writer(response)
    if is_rag:
        writer.writerow(['question_id', 'context', 'question', 'correct_answer', 'subject', 'difficulty'])
        for q in questions:
            writer.writerow([q.question_id, q.context, q.question, q.correct_answer, q.subject, q.difficulty])
    else:
        writer.writerow(['question_id', 'question', 'choice_a', 'choice_b', 'choice_c', 'choice_d',
                         'correct_answer', 'subject', 'difficulty'])
        for q in questions:
            writer.writerow([
                q.question_id, q.question, q.choice_a or '', q.choice_b or '',
                q.choice_c or '', q.choice_d or '', q.correct_answer,
                q.subject, q.difficulty,
            ])
    return response


def export_questions_excel(request, slug):
    benchmark = get_object_or_404(Benchmark, slug=slug)
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        messages.error(request, 'openpyxl is not installed. Please install it with: pip install openpyxl')
        return redirect('benchmarks:detail', slug=slug)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Questions'

    headers = ['question_id', 'question', 'choice_a', 'choice_b', 'choice_c', 'choice_d',
               'correct_answer', 'subject', 'difficulty']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)

    questions = BenchmarkQuestion.objects.filter(benchmark=benchmark)
    for row_idx, q in enumerate(questions, 2):
        ws.cell(row=row_idx, column=1, value=q.question_id)
        ws.cell(row=row_idx, column=2, value=q.question)
        ws.cell(row=row_idx, column=3, value=q.choice_a or '')
        ws.cell(row=row_idx, column=4, value=q.choice_b or '')
        ws.cell(row=row_idx, column=5, value=q.choice_c or '')
        ws.cell(row=row_idx, column=6, value=q.choice_d or '')
        ws.cell(row=row_idx, column=7, value=q.correct_answer)
        ws.cell(row=row_idx, column=8, value=q.subject)
        ws.cell(row=row_idx, column=9, value=q.difficulty)

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{slug}_questions.xlsx"'
    wb.save(response)
    return response


def _run_csv_import_background(run_id, benchmark_id, slug, csv_content):
    """Process a CSV import in a background thread."""
    from apps.runs.models import BenchmarkRun, RunResult
    from apps.benchmarks.models import BenchmarkQuestion

    try:
        run = BenchmarkRun.objects.get(id=run_id)
        benchmark = Benchmark.objects.get(id=benchmark_id)
        loader = get_loader(slug)

        rows = list(csv.DictReader(io.StringIO(csv_content)))
        valid_rows = [r for r in rows if str(r.get('question_id', '')).strip()]

        # Set total up front so the progress bar is meaningful immediately
        run.total_questions = len(valid_rows)
        run.num_questions = len(valid_rows)
        run.save(update_fields=['total_questions', 'num_questions'])

        correct_count = 0
        processed = 0

        for row in valid_rows:
            question_id = str(row.get('question_id', '')).strip()
            answer = (row.get('answer') or '').strip()

            try:
                q = BenchmarkQuestion.objects.get(benchmark=benchmark, question_id=question_id)
            except BenchmarkQuestion.DoesNotExist:
                continue

            if loader:
                is_correct, parsed = loader.evaluate_answer(q, answer)
            else:
                parsed = answer.upper()[:1]
                is_correct = parsed == q.correct_answer.strip().upper()

            RunResult.objects.create(
                run=run,
                question=q,
                model_response=answer,
                parsed_answer=parsed[:100],
                is_correct=is_correct,
                response_time=0.0,
            )
            processed += 1
            if is_correct:
                correct_count += 1

            # Flush progress to DB every 100 rows
            if processed % 100 == 0:
                run.correct_answers = correct_count
                run.score = round(correct_count / processed * 100, 2)
                run.save(update_fields=['correct_answers', 'score'])

        total = run.results.count()  # actual persisted results
        run.total_questions = total
        run.num_questions = total
        run.correct_answers = correct_count
        run.score = round(correct_count / total * 100, 2) if total > 0 else 0.0
        run.status = 'completed'
        run.completed_at = timezone.now()
        run.save()

    except Exception as e:
        try:
            run = BenchmarkRun.objects.get(id=run_id)
            run.status = 'failed'
            run.error_message = str(e)
            run.save()
        except Exception:
            pass


def import_answers_csv(request, slug):
    benchmark = get_object_or_404(Benchmark, slug=slug)
    is_rag = (benchmark.metadata or {}).get('is_rag', False)

    if request.method != 'POST':
        return render(request, 'benchmarks/import_answers.html', {
            'benchmark': benchmark,
            'is_rag': is_rag,
        })

    csv_file = request.FILES.get('csv_file')
    model_name = request.POST.get('model_name', '').strip() or 'CSV / Manual Import'

    if not csv_file:
        messages.error(request, 'Please upload a CSV file.')
        return render(request, 'benchmarks/import_answers.html', {
            'benchmark': benchmark,
            'is_rag': is_rag,
        })

    from apps.providers.models import Provider
    from apps.runs.models import BenchmarkRun

    try:
        csv_provider, _ = Provider.objects.get_or_create(
            slug='csv-manual',
            defaults={
                'name': 'CSV / Manual',
                'provider_type': 'custom',
                'default_model': 'manual',
                'is_active': True,
            }
        )
    except Exception as e:
        messages.error(request, f'Could not create CSV provider: {e}')
        return render(request, 'benchmarks/import_answers.html', {
            'benchmark': benchmark,
            'is_rag': is_rag,
        })

    # Read CSV content before the thread starts (file object is closed after the request)
    try:
        csv_content = csv_file.read().decode('utf-8-sig')
    except Exception as e:
        messages.error(request, f'Could not read CSV file: {e}')
        return render(request, 'benchmarks/import_answers.html', {
            'benchmark': benchmark,
            'is_rag': is_rag,
        })

    # Validate columns before creating any DB records
    try:
        reader = csv.DictReader(io.StringIO(csv_content))
        columns = reader.fieldnames or []
        required = {'question_id', 'answer'}
        missing = required - {c.strip().lower() for c in columns}
        if missing:
            missing_str = ', '.join(f'<code>{c}</code>' for c in sorted(missing))
            found_str = ', '.join(f'<code>{c}</code>' for c in columns) if columns else '(none detected)'
            from django.utils.safestring import mark_safe
            messages.error(request, mark_safe(
                f'Missing required column(s): {missing_str}. '
                f'Columns found in your CSV: {found_str}.'
            ))
            return render(request, 'benchmarks/import_answers.html', {
                'benchmark': benchmark,
                'is_rag': is_rag,
            })
        if not columns:
            messages.error(request, 'The CSV file appears to be empty or has no header row.')
            return render(request, 'benchmarks/import_answers.html', {
                'benchmark': benchmark,
                'is_rag': is_rag,
            })
    except Exception as e:
        messages.error(request, f'Could not parse CSV: {e}')
        return render(request, 'benchmarks/import_answers.html', {
            'benchmark': benchmark,
            'is_rag': is_rag,
        })

    run = BenchmarkRun.objects.create(
        benchmark=benchmark,
        provider=csv_provider,
        model_name=model_name,
        status='running',
        started_at=timezone.now(),
        metadata={'source': 'csv_import', 'is_rag': is_rag},
        tags='csv-import',
    )

    t = threading.Thread(
        target=_run_csv_import_background,
        args=(run.pk, benchmark.pk, slug, csv_content),
        daemon=True,
    )
    t.start()

    messages.info(request, f'CSV import started for "{model_name}" — processing in the background.')
    return redirect('runs:detail', pk=run.pk)


def leaderboard_view(request):
    from apps.runs.models import BenchmarkRun

    benchmark_filter = request.GET.get('benchmark', '')
    show_all = request.GET.get('show_all', '') == '1'

    qs = BenchmarkRun.objects.filter(status='completed').select_related('benchmark', 'provider')

    if benchmark_filter:
        qs = qs.filter(benchmark__slug=benchmark_filter)

    qs = qs.order_by('benchmark__name', '-score', '-created_at')

    # Group by benchmark
    benchmarks_seen = set()
    leaderboard = []
    all_runs = []
    rank_by_benchmark = {}

    for run in qs:
        b_slug = run.benchmark.slug
        all_runs.append(run)
        if b_slug not in rank_by_benchmark:
            rank_by_benchmark[b_slug] = 1
        else:
            rank_by_benchmark[b_slug] += 1

    # Build ranked list
    rank_counter = {}
    for run in all_runs:
        b_slug = run.benchmark.slug
        if b_slug not in rank_counter:
            rank_counter[b_slug] = 1
        else:
            rank_counter[b_slug] += 1
        ci_low, ci_high = _wilson_ci(run.correct_answers, run.total_questions)
        leaderboard.append({
            'run': run,
            'rank': rank_counter[b_slug],
            'ci_low': ci_low,
            'ci_high': ci_high,
        })

    # All benchmarks for filter dropdown
    all_benchmarks = Benchmark.objects.filter(runs__status='completed').distinct().order_by('name')

    return render(request, 'benchmarks/leaderboard.html', {
        'leaderboard': leaderboard,
        'all_benchmarks': all_benchmarks,
        'benchmark_filter': benchmark_filter,
        'show_all': show_all,
    })


def create_custom_benchmark(request):
    if request.method == 'GET':
        return render(request, 'benchmarks/custom_benchmark.html', {
            'categories': Benchmark.CATEGORY_CHOICES,
        })

    name = request.POST.get('name', '').strip()
    description = request.POST.get('description', '').strip()
    category = request.POST.get('category', 'knowledge')
    csv_file = request.FILES.get('csv_file')

    if not name or not csv_file:
        messages.error(request, 'Name and CSV file are required.')
        return render(request, 'benchmarks/custom_benchmark.html', {
            'categories': Benchmark.CATEGORY_CHOICES,
        })

    try:
        import pandas as pd
        content = csv_file.read().decode('utf-8')
        df = pd.read_csv(io.StringIO(content))
    except ImportError:
        messages.error(request, 'pandas is required for CSV import. Install with: pip install pandas')
        return render(request, 'benchmarks/custom_benchmark.html', {
            'categories': Benchmark.CATEGORY_CHOICES,
        })
    except Exception as e:
        messages.error(request, f'Error reading CSV: {e}')
        return render(request, 'benchmarks/custom_benchmark.html', {
            'categories': Benchmark.CATEGORY_CHOICES,
        })

    if 'question' not in df.columns or 'correct_answer' not in df.columns:
        messages.error(request, 'CSV must have at least "question" and "correct_answer" columns.')
        return render(request, 'benchmarks/custom_benchmark.html', {
            'categories': Benchmark.CATEGORY_CHOICES,
        })

    from django.utils.text import slugify
    slug = slugify(name)
    base_slug = slug
    counter = 1
    while Benchmark.objects.filter(slug=slug).exists():
        slug = f"{base_slug}-{counter}"
        counter += 1

    benchmark = Benchmark.objects.create(
        slug=slug,
        name=name,
        description=description,
        category=category,
        metadata={'is_custom': True},
        loaded_at=timezone.now(),
    )

    questions_batch = []
    for idx, row in df.iterrows():
        question_id = f"custom_{idx + 1}"
        questions_batch.append(BenchmarkQuestion(
            benchmark=benchmark,
            question_id=question_id,
            question=str(row.get('question', '')),
            choice_a=str(row['choice_a']) if 'choice_a' in df.columns and pd.notna(row.get('choice_a')) else None,
            choice_b=str(row['choice_b']) if 'choice_b' in df.columns and pd.notna(row.get('choice_b')) else None,
            choice_c=str(row['choice_c']) if 'choice_c' in df.columns and pd.notna(row.get('choice_c')) else None,
            choice_d=str(row['choice_d']) if 'choice_d' in df.columns and pd.notna(row.get('choice_d')) else None,
            correct_answer=str(row.get('correct_answer', '')),
            subject=str(row['subject']) if 'subject' in df.columns and pd.notna(row.get('subject')) else '',
            difficulty=str(row['difficulty']) if 'difficulty' in df.columns and pd.notna(row.get('difficulty')) else '',
        ))

    BenchmarkQuestion.objects.bulk_create(questions_batch)
    benchmark.num_questions = len(questions_batch)
    benchmark.save()

    messages.success(request, f'Custom benchmark "{name}" created with {len(questions_batch)} questions.')
    return redirect('benchmarks:detail', slug=benchmark.slug)


def download_example_csv(request):
    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = 'attachment; filename="example_benchmark.csv"'
    writer = csv.writer(response)
    writer.writerow(['question', 'choice_a', 'choice_b', 'choice_c', 'choice_d', 'correct_answer', 'subject', 'difficulty'])
    writer.writerow([
        'What is the capital of France?',
        'London', 'Paris', 'Berlin', 'Madrid',
        'B', 'Geography', 'easy'
    ])
    writer.writerow([
        'Which planet is known as the Red Planet?',
        'Venus', 'Jupiter', 'Mars', 'Saturn',
        'C', 'Science', 'easy'
    ])
    return response


def hf_import_view(request):
    if request.method == 'GET':
        return render(request, 'benchmarks/hf_import.html', {
            'categories': Benchmark.CATEGORY_CHOICES,
        })

    dataset_path = request.POST.get('dataset_path', '').strip()
    config_name = request.POST.get('config_name', '').strip() or None
    split = request.POST.get('split', 'test').strip() or 'test'
    question_col = request.POST.get('question_col', 'question').strip()
    choice_a_col = request.POST.get('choice_a_col', '').strip()
    choice_b_col = request.POST.get('choice_b_col', '').strip()
    choice_c_col = request.POST.get('choice_c_col', '').strip()
    choice_d_col = request.POST.get('choice_d_col', '').strip()
    answer_col = request.POST.get('answer_col', 'answer').strip()
    subject_col = request.POST.get('subject_col', '').strip() or None
    benchmark_name = request.POST.get('benchmark_name', '').strip()
    benchmark_slug = request.POST.get('benchmark_slug', '').strip()
    category = request.POST.get('category', 'knowledge')
    num_samples = int(request.POST.get('num_samples', 0) or 0)

    if not dataset_path or not benchmark_name:
        messages.error(request, 'Dataset path and benchmark name are required.')
        return render(request, 'benchmarks/hf_import.html', {
            'categories': Benchmark.CATEGORY_CHOICES,
        })

    try:
        from datasets import load_dataset
        dataset = load_dataset(dataset_path, config_name, split=split)
    except Exception as e:
        messages.error(request, f'Error loading HuggingFace dataset: {e}')
        return render(request, 'benchmarks/hf_import.html', {
            'categories': Benchmark.CATEGORY_CHOICES,
        })

    from django.utils.text import slugify
    if not benchmark_slug:
        benchmark_slug = slugify(benchmark_name)

    base_slug = benchmark_slug
    counter = 1
    while Benchmark.objects.filter(slug=benchmark_slug).exists():
        benchmark_slug = f"{base_slug}-{counter}"
        counter += 1

    benchmark = Benchmark.objects.create(
        slug=benchmark_slug,
        name=benchmark_name,
        category=category,
        metadata={'source': dataset_path, 'config': config_name, 'split': split},
        loaded_at=timezone.now(),
    )

    questions_batch = []
    items = list(dataset)
    if num_samples and num_samples > 0:
        items = items[:num_samples]

    for idx, item in enumerate(items):
        def get_col(col):
            if col and col in item:
                return str(item[col])
            return None

        questions_batch.append(BenchmarkQuestion(
            benchmark=benchmark,
            question_id=str(idx + 1),
            question=str(item.get(question_col, '')),
            choice_a=get_col(choice_a_col),
            choice_b=get_col(choice_b_col),
            choice_c=get_col(choice_c_col),
            choice_d=get_col(choice_d_col),
            correct_answer=str(item.get(answer_col, '')),
            subject=str(item[subject_col]) if subject_col and subject_col in item else '',
        ))

    BenchmarkQuestion.objects.bulk_create(questions_batch)
    benchmark.num_questions = len(questions_batch)
    benchmark.save()

    messages.success(request, f'Imported {len(questions_batch)} questions from "{dataset_path}" as "{benchmark_name}".')
    return redirect('benchmarks:detail', slug=benchmark.slug)


def hardness_tracker_view(request):
    from apps.runs.models import RunResult

    benchmark_filter = request.GET.get('benchmark', '')
    min_attempts = int(request.GET.get('min_attempts', 2) or 2)

    questions_qs = BenchmarkQuestion.objects.annotate(
        total_attempts=Count('results'),
        correct_attempts=Count('results', filter=Q(results__is_correct=True)),
    ).filter(total_attempts__gte=min_attempts)

    if benchmark_filter:
        questions_qs = questions_qs.filter(benchmark__slug=benchmark_filter)

    questions_with_fail_rate = []
    for q in questions_qs.select_related('benchmark'):
        total = q.total_attempts
        correct = q.correct_attempts
        fail_rate = round(1.0 - (correct / total), 3) if total > 0 else 0.0
        questions_with_fail_rate.append({
            'question': q,
            'total': total,
            'correct': correct,
            'fail_rate': fail_rate,
            'fail_pct': round(fail_rate * 100, 1),
        })

    questions_with_fail_rate.sort(key=lambda x: x['fail_rate'], reverse=True)

    hard_questions = [q for q in questions_with_fail_rate if q['fail_rate'] > 0.8]
    controversial_questions = [q for q in questions_with_fail_rate if 0.3 < q['fail_rate'] <= 0.8]
    easy_questions = [q for q in questions_with_fail_rate if q['fail_rate'] <= 0.3]

    all_benchmarks = Benchmark.objects.filter(questions__results__isnull=False).distinct()

    return render(request, 'benchmarks/hardness.html', {
        'hard_questions': hard_questions[:50],
        'controversial_questions': controversial_questions[:50],
        'easy_questions': easy_questions[:20],
        'all_questions': questions_with_fail_rate[:100],
        'all_benchmarks': all_benchmarks,
        'benchmark_filter': benchmark_filter,
        'min_attempts': min_attempts,
        'total_analyzed': len(questions_with_fail_rate),
    })


# Prompt Library Views

def prompt_list_view(request):
    prompts = PromptTemplate.objects.all()
    return render(request, 'prompts/list.html', {'prompts': prompts})


def prompt_create_view(request):
    if request.method == 'POST':
        name = request.POST.get('name', '').strip()
        description = request.POST.get('description', '').strip()
        content = request.POST.get('content', '').strip()
        is_default = request.POST.get('is_default', '') == 'on'

        if not name or not content:
            messages.error(request, 'Name and content are required.')
            return render(request, 'prompts/create.html', {
                'name': name,
                'description': description,
                'content': content,
            })

        prompt = PromptTemplate.objects.create(
            name=name,
            description=description,
            content=content,
            is_default=is_default,
        )
        messages.success(request, f'Prompt template "{name}" created.')
        return redirect('benchmarks:prompt_list')

    return render(request, 'prompts/create.html', {})


def prompt_edit_view(request, pk):
    prompt = get_object_or_404(PromptTemplate, pk=pk)
    if request.method == 'POST':
        prompt.name = request.POST.get('name', '').strip()
        prompt.description = request.POST.get('description', '').strip()
        prompt.content = request.POST.get('content', '').strip()
        prompt.is_default = request.POST.get('is_default', '') == 'on'

        if not prompt.name or not prompt.content:
            messages.error(request, 'Name and content are required.')
            return render(request, 'prompts/create.html', {'prompt': prompt})

        prompt.save()
        messages.success(request, f'Prompt template "{prompt.name}" updated.')
        return redirect('benchmarks:prompt_list')

    return render(request, 'prompts/create.html', {'prompt': prompt})


def prompt_delete_view(request, pk):
    prompt = get_object_or_404(PromptTemplate, pk=pk)
    if request.method == 'POST':
        name = prompt.name
        prompt.delete()
        messages.success(request, f'Prompt template "{name}" deleted.')
        return redirect('benchmarks:prompt_list')
    return render(request, 'prompts/delete_confirm.html', {'prompt': prompt})


def prompt_api_view(request):
    """API endpoint returning prompts as JSON for modal."""
    prompts = list(PromptTemplate.objects.values('id', 'name', 'description', 'content'))
    return JsonResponse({'prompts': prompts})


# ---------------------------------------------------------------------------
# RAG Benchmark from CSV
# ---------------------------------------------------------------------------

def create_rag_benchmark(request):
    """
    Create a RAG (Retrieval-Augmented Generation) benchmark from a CSV file.

    Required CSV columns:
        context         — the retrieved passage/document
        question        — the question to answer using the context
        correct_answer  — expected answer; use | to separate aliases, e.g. "Paris|City of Light"

    Optional columns:
        subject, difficulty

    The benchmark is tagged with metadata={'is_rag': True} so the RAGBenchmarkLoader
    is automatically selected during evaluation runs.
    """
    if request.method == 'GET':
        return render(request, 'benchmarks/rag_benchmark.html', {
            'categories': Benchmark.CATEGORY_CHOICES,
        })

    name = request.POST.get('name', '').strip()
    description = request.POST.get('description', '').strip()
    category = request.POST.get('category', 'knowledge')
    csv_file = request.FILES.get('csv_file')

    if not name or not csv_file:
        messages.error(request, 'Name and CSV file are required.')
        return render(request, 'benchmarks/rag_benchmark.html', {'categories': Benchmark.CATEGORY_CHOICES})

    try:
        import pandas as pd
        content = csv_file.read().decode('utf-8')
        df = pd.read_csv(io.StringIO(content))
    except Exception as e:
        messages.error(request, f'Error reading CSV: {e}')
        return render(request, 'benchmarks/rag_benchmark.html', {'categories': Benchmark.CATEGORY_CHOICES})

    required_cols = {'context', 'question', 'correct_answer'}
    missing = required_cols - set(df.columns)
    if missing:
        messages.error(request, f'CSV is missing required columns: {", ".join(sorted(missing))}')
        return render(request, 'benchmarks/rag_benchmark.html', {'categories': Benchmark.CATEGORY_CHOICES})

    # Remove rows with empty context or question
    df = df.dropna(subset=['context', 'question', 'correct_answer'])
    if df.empty:
        messages.error(request, 'No valid rows found after removing empty context/question/answer rows.')
        return render(request, 'benchmarks/rag_benchmark.html', {'categories': Benchmark.CATEGORY_CHOICES})

    from django.utils.text import slugify
    slug = slugify(name)
    base_slug, counter = slug, 1
    while Benchmark.objects.filter(slug=slug).exists():
        slug = f"{base_slug}-{counter}"
        counter += 1

    benchmark = Benchmark.objects.create(
        slug=slug,
        name=name,
        description=description,
        category=category,
        metadata={'is_rag': True, 'is_custom': True},
        loaded_at=timezone.now(),
    )

    questions_batch = []
    for idx, row in df.iterrows():
        question_id = f"rag_{idx + 1}"
        # Support aliases in correct_answer via | separator (stored as-is, evaluated by RAGBenchmarkLoader)
        correct_answer = str(row.get('correct_answer', '')).strip()
        questions_batch.append(BenchmarkQuestion(
            benchmark=benchmark,
            question_id=question_id,
            question=str(row.get('question', '')).strip(),
            context=str(row.get('context', '')).strip(),
            correct_answer=correct_answer,
            subject=str(row['subject']).strip() if 'subject' in df.columns and pd.notna(row.get('subject')) else '',
            difficulty=str(row['difficulty']).strip() if 'difficulty' in df.columns and pd.notna(row.get('difficulty')) else '',
        ))

    BenchmarkQuestion.objects.bulk_create(questions_batch)
    benchmark.num_questions = len(questions_batch)
    benchmark.save()

    messages.success(request, f'RAG benchmark "{name}" created with {len(questions_batch)} questions.')
    return redirect('benchmarks:detail', slug=benchmark.slug)


def set_prompt_template_view(request, slug):
    if request.method != 'POST':
        return redirect('benchmarks:detail', slug=slug)
    benchmark = get_object_or_404(Benchmark, slug=slug)
    benchmark.prompt_template = request.POST.get('prompt_template', '').strip()
    benchmark.save(update_fields=['prompt_template'])
    messages.success(request, 'Prompt template saved.')
    return redirect('benchmarks:detail', slug=slug)


def download_rag_example_csv(request):
    """Download a sample CSV demonstrating the RAG benchmark format."""
    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = 'attachment; filename="rag_benchmark_example.csv"'
    writer = csv.writer(response)
    writer.writerow(['context', 'question', 'correct_answer', 'subject', 'difficulty'])
    writer.writerow([
        'The mitochondria is a membrane-bound organelle found in the cytoplasm of eukaryotic cells. '
        'It generates most of the cell\'s supply of adenosine triphosphate (ATP), used as a source of chemical energy.',
        'What is the primary function of the mitochondria?',
        'produce ATP|generate energy|energy production',
        'cell biology',
        'easy',
    ])
    writer.writerow([
        'CRISPR-Cas9 is a molecular tool derived from a bacterial immune system. '
        'It allows scientists to edit sections of the genome by removing, adding or altering sections of DNA sequence. '
        'The system uses a guide RNA to direct the Cas9 protein to the target DNA.',
        'What molecule guides the Cas9 protein to its target DNA sequence?',
        'guide RNA|gRNA',
        'genomics',
        'medium',
    ])
    writer.writerow([
        'The Hardy-Weinberg principle states that allele and genotype frequencies in a population '
        'will remain constant from generation to generation in the absence of evolutionary influences. '
        'These influences include mate choice, mutation, selection, genetic drift and gene flow.',
        'According to Hardy-Weinberg, what must be absent for allele frequencies to remain constant?',
        'evolutionary influences|natural selection|mutation and selection',
        'population genetics',
        'hard',
    ])
    return response
