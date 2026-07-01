"""
Export run results to CSV or Excel.

Examples:
    python manage.py bm_export 42
    python manage.py bm_export 42 --format excel --output results_42.xlsx
    python manage.py bm_export 42 --format csv --output /tmp/run42.csv
"""
import os

from django.core.management.base import BaseCommand, CommandError

from apps.runs.models import BenchmarkRun


class Command(BaseCommand):
    help = 'Export run results to CSV or Excel'

    def add_arguments(self, parser):
        parser.add_argument('run_id', type=int, help='Run ID to export')
        parser.add_argument('--format', choices=['csv', 'excel'], default='csv',
                            help='Output format (default: csv)')
        parser.add_argument('--output', '-o', default='',
                            help='Output file path (default: run_<id>.<ext>)')

    def handle(self, *args, **options):
        try:
            run = BenchmarkRun.objects.get(id=options['run_id'])
        except BenchmarkRun.DoesNotExist:
            raise CommandError(f"Run #{options['run_id']} not found.")

        fmt = options['format']
        ext = 'xlsx' if fmt == 'excel' else 'csv'
        output = options['output'] or f'run_{run.id}.{ext}'

        if fmt == 'csv':
            self._export_csv(run, output)
        else:
            self._export_excel(run, output)

        self.stdout.write(self.style.SUCCESS(f'Exported to: {os.path.abspath(output)}'))

    def _export_csv(self, run, path):
        import csv
        results = run.results.select_related('question').order_by('id')
        with open(path, 'w', newline='', encoding='utf-8') as f:
            writer = csv.writer(f)
            writer.writerow([
                'question_id', 'subject', 'difficulty',
                'question', 'correct_answer',
                'model_answer', 'parsed_answer', 'is_correct',
                'response_time_s', 'tokens_input', 'tokens_output', 'cost_usd',
            ])
            for r in results:
                writer.writerow([
                    r.question.question_id,
                    r.question.subject,
                    r.question.difficulty,
                    r.question.question,
                    r.question.correct_answer,
                    r.model_response,
                    r.parsed_answer,
                    r.is_correct,
                    round(r.response_time or 0, 4),
                    r.tokens_input or '',
                    r.tokens_output or '',
                    round(r.estimated_cost or 0, 6),
                ])

    def _export_excel(self, run, path):
        try:
            import openpyxl
            from openpyxl.styles import Font
        except ImportError:
            raise CommandError(
                'openpyxl is required: pip install openpyxl'
            )

        wb = openpyxl.Workbook()

        # Sheet 1 — Summary
        ws_sum = wb.active
        ws_sum.title = 'Summary'
        rows = [
            ('Run ID', run.id),
            ('Benchmark', run.benchmark.name),
            ('Provider', run.provider.name),
            ('Model', run.model_name),
            ('Status', run.status),
            ('Score', f'{run.score:.1f}%'),
            ('Correct', run.correct_answers),
            ('Total', run.total_questions),
            ('Temperature', run.temperature),
            ('Max Tokens', run.max_tokens),
            ('Created', str(run.created_at)),
            ('Completed', str(run.completed_at or '-')),
            ('Duration (s)', str(run.duration_seconds or '-')),
            ('Tags', run.tags or ''),
            ('Notes', run.notes or ''),
        ]
        for row_idx, (label, value) in enumerate(rows, 1):
            ws_sum.cell(row=row_idx, column=1, value=label).font = Font(bold=True)
            ws_sum.cell(row=row_idx, column=2, value=value)

        # Sheet 2 — By Subject
        from django.db.models import Count, Avg
        ws_sub = wb.create_sheet('By Subject')
        ws_sub.append(['Subject', 'Correct', 'Total', 'Score %'])
        for cell in ws_sub[1]:
            cell.font = Font(bold=True)
        subjects = (
            run.results.values('question__subject')
            .annotate(
                total=Count('id'),
                correct=Count('id', filter=__import__('django.db.models', fromlist=['Q']).Q(is_correct=True)),
            )
            .order_by('question__subject')
        )
        for s in subjects:
            pct = round(s['correct'] / s['total'] * 100, 1) if s['total'] else 0
            ws_sub.append([s['question__subject'], s['correct'], s['total'], pct])

        # Sheet 3 — Results
        ws_res = wb.create_sheet('Results')
        headers = ['question_id', 'subject', 'difficulty', 'question',
                   'correct_answer', 'model_answer', 'parsed_answer',
                   'is_correct', 'response_time_s', 'tokens_input', 'tokens_output']
        ws_res.append(headers)
        for cell in ws_res[1]:
            cell.font = Font(bold=True)
        for r in run.results.select_related('question').order_by('id'):
            ws_res.append([
                r.question.question_id,
                r.question.subject,
                r.question.difficulty,
                r.question.question,
                r.question.correct_answer,
                r.model_response,
                r.parsed_answer,
                r.is_correct,
                round(r.response_time or 0, 4),
                r.tokens_input or 0,
                r.tokens_output or 0,
            ])

        wb.save(path)
